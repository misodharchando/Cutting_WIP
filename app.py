from __future__ import annotations

from io import BytesIO
from pathlib import Path
from copy import copy
import re

import numpy as np
import pandas as pd
import plotly.express as px
import streamlit as st

st.set_page_config(page_title="Cutting WIP Control Tower", page_icon="🏭", layout="wide")

REQUIRED_COLUMNS = [
    "Customer", "Style", "Season", "Sales Order", "Order Qty", "Cut Qty",
    "Balance to Cut Qty", "Bundled Qty", "Issued Qty", "Balance to Issue Qty",
]
NUMERIC_COLUMNS = [
    "Order Qty", "Cut Qty", "Balance to Cut Qty", "Bundled Qty",
    "Issued Qty", "Balance to Issue Qty",
]

ALIASES = {
    "customer": "Customer", "style": "Style", "season": "Season",
    "sales order": "Sales Order", "salesorder": "Sales Order", "so": "Sales Order",
    "order qty": "Order Qty", "order quantity": "Order Qty",
    "cut qty": "Cut Qty", "cut quantity": "Cut Qty",
    "balance to cut qty": "Balance to Cut Qty", "balance to cut": "Balance to Cut Qty",
    "bundled qty": "Bundled Qty", "bundle qty": "Bundled Qty",
    "issued qty": "Issued Qty", "issue qty": "Issued Qty",
    "balance to issue qty": "Balance to Issue Qty", "balance to issue": "Balance to Issue Qty",
}


def clean_header(value: object) -> str:
    text = re.sub(r"\s+", " ", str(value).replace("\n", " ").strip())
    return ALIASES.get(text.lower(), text)


def locate_header(raw: pd.DataFrame) -> int:
    expected = {x.lower() for x in REQUIRED_COLUMNS}
    best_row, best_score = 0, -1
    for idx in range(min(30, len(raw))):
        values = {clean_header(x).lower() for x in raw.iloc[idx].tolist() if pd.notna(x)}
        score = len(values & expected)
        if score > best_score:
            best_row, best_score = idx, score
    if best_score < 5:
        raise ValueError("Could not detect the header row. Please upload the standard Cutting WIP Inquiry export.")
    return best_row


@st.cache_data(show_spinner=False)
def load_workbook(file_bytes: bytes, sheet_name: str | int = 0) -> pd.DataFrame:
    try:
        # Calamine reads cell values without depending on potentially damaged Excel style records.
        raw = pd.read_excel(BytesIO(file_bytes), sheet_name=sheet_name, header=None, engine="calamine")
    except Exception:
        # Fallback for environments where calamine is unavailable.
        raw = pd.read_excel(BytesIO(file_bytes), sheet_name=sheet_name, header=None, engine="openpyxl")
    header_row = locate_header(raw)
    headers = [clean_header(x) for x in raw.iloc[header_row].tolist()]
    df = raw.iloc[header_row + 1:].copy()
    df.columns = headers
    df = df.loc[:, ~pd.Index(df.columns).duplicated()]

    missing = [c for c in REQUIRED_COLUMNS if c not in df.columns]
    if missing:
        raise ValueError("Missing required columns: " + ", ".join(missing))

    df = df[REQUIRED_COLUMNS].copy()
    df = df.dropna(how="all")
    for c in ["Customer", "Style", "Season", "Sales Order"]:
        df[c] = df[c].fillna("").astype(str).str.strip()

    # Remove report totals and blank rows. Totals are recalculated from order-level records.
    df = df[(df["Sales Order"].str.lower() != "total") & (df["Customer"].str.lower() != "total")]
    df = df[df[["Customer", "Style", "Season", "Sales Order"]].ne("").any(axis=1)]

    for c in NUMERIC_COLUMNS:
        df[c] = pd.to_numeric(df[c], errors="coerce").fillna(0)

    df = df[df[NUMERIC_COLUMNS].abs().sum(axis=1) > 0].reset_index(drop=True)
    return enrich(df)


def enrich(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    out["Calculated Balance to Cut"] = (out["Order Qty"] - out["Cut Qty"]).clip(lower=0)
    out["Cut to Bundle WIP"] = (out["Cut Qty"] - out["Bundled Qty"]).clip(lower=0)
    out["Bundle to Issue WIP"] = (out["Bundled Qty"] - out["Issued Qty"]).clip(lower=0)
    out["Cut Completion %"] = np.where(out["Order Qty"] > 0, out["Cut Qty"] / out["Order Qty"], 0)
    out["Bundle Completion %"] = np.where(out["Cut Qty"] > 0, out["Bundled Qty"] / out["Cut Qty"], 0)
    out["Issue Completion %"] = np.where(out["Bundled Qty"] > 0, out["Issued Qty"] / out["Bundled Qty"], 0)
    out["Balance to Cut Variance"] = out["Balance to Cut Qty"] - out["Calculated Balance to Cut"]
    out["Overcut Qty"] = (out["Cut Qty"] - out["Order Qty"]).clip(lower=0)
    out["Over-bundled Qty"] = (out["Bundled Qty"] - out["Cut Qty"]).clip(lower=0)
    out["Over-issued Qty"] = (out["Issued Qty"] - out["Bundled Qty"]).clip(lower=0)

    conditions = [
        (out["Balance to Cut Qty"] >= 5000) | (out["Balance to Issue Qty"] >= 5000),
        (out["Balance to Cut Qty"] >= 1000) | (out["Balance to Issue Qty"] >= 1000),
    ]
    out["Priority"] = np.select(conditions, ["Red", "Yellow"], default="Green")

    flags = []
    for _, r in out.iterrows():
        row_flags = []
        if r["Overcut Qty"] > 0: row_flags.append("Overcut")
        if r["Over-bundled Qty"] > 0: row_flags.append("Over-bundled")
        if r["Over-issued Qty"] > 0: row_flags.append("Over-issued")
        if abs(r["Balance to Cut Variance"]) > 0.01: row_flags.append("Balance mismatch")
        if r["Cut Qty"] > 0 and r["Issued Qty"] == 0: row_flags.append("Cut but not issued")
        flags.append(", ".join(row_flags) if row_flags else "OK")
    out["Exception"] = flags
    return out


def fmt_number(value: float) -> str:
    return f"{value:,.0f}"


def pct(num: float, den: float) -> float:
    return (num / den) if den else 0.0


def to_excel(data: pd.DataFrame, summary: pd.DataFrame) -> bytes:
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        data.to_excel(writer, index=False, sheet_name="Filtered WIP")
        summary.to_excel(writer, index=False, sheet_name="Customer Summary")
        for ws in writer.book.worksheets:
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                from openpyxl.styles import Font, PatternFill
                cell.font = Font(name=cell.font.name or "Calibri", size=cell.font.sz or 11, bold=True, color="FFFFFF")
                cell.fill = PatternFill(fill_type="solid", fgColor="1F4E78")
            for col in ws.columns:
                width = min(max(len(str(c.value or "")) for c in col) + 2, 45)
                ws.column_dimensions[col[0].column_letter].width = width
    return output.getvalue()


st.markdown("""
<style>
.block-container {padding-top: 1.25rem; padding-bottom: 2rem;}
div[data-testid="stMetric"] {background:#f7fafc; border:1px solid #dbe5ee; padding:12px; border-radius:12px;}
.red-box {background:#ffe9e9; border-left:5px solid #d62828; padding:10px 14px; border-radius:8px;}
</style>
""", unsafe_allow_html=True)

st.title("🏭 Cutting WIP Control Tower")
st.caption("Quantity-based cutting, bundling and issue monitoring from the HMES Cutting WIP Inquiry export")

with st.sidebar:
    st.header("Data source")
    uploaded = st.file_uploader("Upload Cutting WIP Inquiry", type=["xlsx", "xls"])
    default_path = Path(__file__).with_name("Cutting WIP Inquiry.xlsx")
    if uploaded is not None:
        file_bytes = uploaded.getvalue()
        source_name = uploaded.name
    elif default_path.exists():
        file_bytes = default_path.read_bytes()
        source_name = default_path.name
        st.success("Sample workbook loaded")
    else:
        file_bytes = None
        source_name = ""

if not file_bytes:
    st.info("Upload the HMES Cutting WIP Inquiry Excel file to begin.")
    st.stop()

try:
    df = load_workbook(file_bytes)
except Exception as exc:
    st.error(f"Unable to read the workbook: {exc}")
    st.stop()

with st.sidebar:
    st.divider()
    st.header("Filters")
    customer_values = sorted(x for x in df["Customer"].unique() if x)
    season_values = sorted(x for x in df["Season"].unique() if x)
    priority_values = ["Red", "Yellow", "Green"]
    customers = st.multiselect("Customer", customer_values)
    seasons = st.multiselect("Season", season_values)
    priorities = st.multiselect("Priority", priority_values)
    search = st.text_input("Search Sales Order / Style")
    exception_only = st.checkbox("Show exceptions only")

filtered = df.copy()
if customers: filtered = filtered[filtered["Customer"].isin(customers)]
if seasons: filtered = filtered[filtered["Season"].isin(seasons)]
if priorities: filtered = filtered[filtered["Priority"].isin(priorities)]
if search:
    mask = filtered["Sales Order"].str.contains(search, case=False, na=False) | filtered["Style"].str.contains(search, case=False, na=False)
    filtered = filtered[mask]
if exception_only: filtered = filtered[filtered["Exception"] != "OK"]

order_qty = filtered["Order Qty"].sum()
cut_qty = filtered["Cut Qty"].sum()
bundled_qty = filtered["Bundled Qty"].sum()
issued_qty = filtered["Issued Qty"].sum()
bal_cut = filtered["Balance to Cut Qty"].sum()
bal_issue = filtered["Balance to Issue Qty"].sum()

st.caption(f"Source: {source_name} | {len(filtered):,} sales-order records shown")

r1 = st.columns(6)
r1[0].metric("Order Qty", fmt_number(order_qty))
r1[1].metric("Cut Qty", fmt_number(cut_qty), f"{pct(cut_qty, order_qty):.1%} of order")
r1[2].metric("Balance to Cut", fmt_number(bal_cut))
r1[3].metric("Bundled Qty", fmt_number(bundled_qty), f"{pct(bundled_qty, cut_qty):.1%} of cut")
r1[4].metric("Issued Qty", fmt_number(issued_qty), f"{pct(issued_qty, bundled_qty):.1%} of bundled")
r1[5].metric("Balance to Issue", fmt_number(bal_issue))

red_count = int((filtered["Priority"] == "Red").sum())
exception_count = int((filtered["Exception"] != "OK").sum())
if red_count:
    st.markdown(f'<div class="red-box"><b>{red_count:,} critical orders</b> meet the current Red threshold. {exception_count:,} rows contain one or more data/process exceptions.</div>', unsafe_allow_html=True)

summary = filtered.groupby("Customer", as_index=False).agg({
    "Order Qty":"sum", "Cut Qty":"sum", "Balance to Cut Qty":"sum",
    "Bundled Qty":"sum", "Issued Qty":"sum", "Balance to Issue Qty":"sum",
}).sort_values("Balance to Issue Qty", ascending=False)

page1, page2, page3, page4 = st.tabs(["Overview", "Priority Board", "Customer & Season", "Data Quality"])

with page1:
    c1, c2 = st.columns(2)
    funnel = pd.DataFrame({"Stage":["Order", "Cut", "Bundled", "Issued"], "Quantity":[order_qty, cut_qty, bundled_qty, issued_qty]})
    c1.plotly_chart(px.funnel(funnel, x="Quantity", y="Stage", title="Order-to-Issue Funnel", color="Stage"), use_container_width=True)
    priority_summary = filtered.groupby("Priority", as_index=False)["Sales Order"].count().rename(columns={"Sales Order":"Orders"})
    priority_summary["Priority"] = pd.Categorical(priority_summary["Priority"], ["Red","Yellow","Green"], ordered=True)
    priority_summary = priority_summary.sort_values("Priority")
    c2.plotly_chart(px.bar(priority_summary, x="Priority", y="Orders", color="Priority", title="Orders by Priority", color_discrete_map={"Red":"#d62828","Yellow":"#f4b400","Green":"#2a9d8f"}), use_container_width=True)

    top = filtered.nlargest(15, "Balance to Issue Qty")
    st.plotly_chart(px.bar(top, x="Balance to Issue Qty", y="Sales Order", color="Customer", orientation="h", title="Top 15 Balance-to-Issue Orders", hover_data=["Style", "Season"]), use_container_width=True)

with page2:
    display_cols = ["Priority", "Customer", "Style", "Season", "Sales Order", "Order Qty", "Cut Qty", "Balance to Cut Qty", "Bundled Qty", "Issued Qty", "Balance to Issue Qty", "Cut Completion %", "Issue Completion %", "Exception"]
    board = filtered.sort_values(["Priority", "Balance to Issue Qty", "Balance to Cut Qty"], ascending=[True, False, False])[display_cols]
    def priority_style(value):
        if value == "Red":
            return "background-color:#ffd6d6;color:#8b0000;font-weight:bold"
        if value == "Yellow":
            return "background-color:#fff1b8;color:#6b5200;font-weight:bold"
        return "background-color:#d9f2e6;color:#075e3a;font-weight:bold"

    styled_board = board.style.format({
        "Cut Completion %": "{:.1%}",
        "Issue Completion %": "{:.1%}",
    })
    # pandas 2.1+ uses Styler.map; older versions use Styler.applymap.
    if hasattr(styled_board, "map"):
        styled_board = styled_board.map(priority_style, subset=["Priority"])
    else:
        styled_board = styled_board.applymap(priority_style, subset=["Priority"])
    st.dataframe(styled_board, use_container_width=True, height=560)
    st.download_button("Download filtered WIP report", to_excel(filtered, summary), file_name="Cutting_WIP_Control_Report.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

with page3:
    left, right = st.columns(2)
    left.plotly_chart(px.bar(summary.head(15), x="Balance to Issue Qty", y="Customer", orientation="h", title="Customer-wise Balance to Issue"), use_container_width=True)
    season_summary = filtered.groupby("Season", as_index=False)[["Order Qty", "Balance to Cut Qty", "Balance to Issue Qty"]].sum().sort_values("Balance to Issue Qty", ascending=False)
    right.plotly_chart(px.bar(season_summary, x="Season", y=["Balance to Cut Qty", "Balance to Issue Qty"], barmode="group", title="Season-wise Pending Quantity"), use_container_width=True)
    st.subheader("Customer Summary")
    st.dataframe(summary, use_container_width=True, hide_index=True)

with page4:
    ex = filtered[filtered["Exception"] != "OK"].copy()
    e1, e2, e3, e4 = st.columns(4)
    e1.metric("Exception Rows", f"{len(ex):,}")
    e2.metric("Overcut Qty", fmt_number(filtered["Overcut Qty"].sum()))
    e3.metric("Over-bundled Qty", fmt_number(filtered["Over-bundled Qty"].sum()))
    e4.metric("Over-issued Qty", fmt_number(filtered["Over-issued Qty"].sum()))
    st.warning("The report's Balance to Issue field is retained as supplied. The app does not replace it with a guessed formula.")
    st.dataframe(ex[["Customer", "Style", "Sales Order", "Order Qty", "Cut Qty", "Bundled Qty", "Issued Qty", "Balance to Cut Qty", "Calculated Balance to Cut", "Balance to Cut Variance", "Exception"]], use_container_width=True, hide_index=True, height=540)

st.divider()
st.caption("Priority thresholds are configurable business assumptions: Red ≥ 5,000 pending pieces; Yellow ≥ 1,000; otherwise Green. This dataset has no transaction timestamps, so the app provides quantity-based WIP, not hourly aging.")
